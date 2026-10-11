"""Timetable import — .xlsx / .csv / .json → TimetableIn (SPEC §40, §46, §47.3, §47.4).

- .xlsx opened read_only + data_only (formulas never evaluated); .xlsm refused; ≤ 1 MB, ≤ 500 rows, ≤ 20 columns
- A sheet qualifies when its first row carries the IR working-timetable headers (by text, case-insensitive,
  never by coordinate); several qualifying sheets → {"needs_sheet": [...]} for the controller to pick
- Station settings read from the "Station" sheet by row label ("Minimum platform clearance (min)",
  "Maximum hold (min)") or the named cells MinClearance / MaxHold
- Time cells: Excel time, Excel datetime (time part) or HH:MM text; anything else is a row error
- Allowed PFs: comma/space list of PF numbers; blank → derived from length (coaches) or goods loop (wagons) and use
- Errors are returned per row and column and nothing is imported until every row passes
- Your own station (§17 load `{station, timetable}`): the "Station" sheet's platform table (PF · Length (coaches) ·
  Goods loop (wagons) · Use · Names) or a .json `{"station": …, "timetable": …}`. A table identical to the loaded
  station keeps it as is; platforms without names get the standard EN · हिं · मर aliases so the parser resolves them
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import re
from typing import Any, Optional

from pydantic import ValidationError

from .models import Platform, Station, TimetableIn, TrainIn, referential_errors

MAX_BYTES, MAX_ROWS, MAX_COLS = 1024 * 1024, 500, 20
REQUIRED = ["train no.", "train name", "type", "from", "to", "dir", "arr", "dep",
            "sch. pf", "req. pf", "allowed pfs", "priority"]
OPTIONAL = ["halt (min)", "coaches", "wagons", "days of run"]
SETTING_LABELS = {"minimum platform clearance (min)": "min_gap", "maximum hold (min)": "max_hold"}
NAMED = {"MinClearance": "min_gap", "MaxHold": "max_hold"}
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


class ImportErrorList(Exception):
    def __init__(self, errors: list[str]):
        super().__init__("; ".join(errors))
        self.errors = errors


def _norm(v: Any) -> str:
    return " ".join(str(v or "").strip().lower().split())


def _header_map(row: tuple) -> Optional[dict[str, int]]:
    names = [_norm(c) for c in row]
    if not all(h in names for h in REQUIRED):
        return None
    return {h: names.index(h) for h in REQUIRED + OPTIONAL if h in names}


def _hhmm(v: Any, where: str) -> str:
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, (int, float)) and 0 <= float(v) < 1:            # raw Excel day fraction
        m = int(round(float(v) * 1440))
        return f"{m // 60:02d}:{m % 60:02d}"
    s = str(v or "").strip()
    if re.fullmatch(r"\d{1,2}:\d{2}(:\d{2})?", s):
        h, mi = s.split(":")[:2]
        if int(h) <= 23 and int(mi) <= 59:
            return f"{int(h):02d}:{mi}"
    raise ValueError(f"{where}: expected a time like 14:05")


def _pf(v: Any, where: str) -> Optional[str]:
    if v is None or str(v).strip() == "":
        return None
    s = str(v).strip().upper().removeprefix("PF").removeprefix("P").strip()
    if not re.fullmatch(r"\d{1,2}(\.0)?", s):
        raise ValueError(f"{where}: expected a PF number like 3")
    return f"P{int(float(s))}"


def _int(v: Any, where: str) -> Optional[int]:
    if v is None or str(v).strip() == "":
        return None
    try:
        return int(float(str(v).strip()))
    except ValueError:
        raise ValueError(f"{where}: expected a whole number") from None


def derive_allowed(station: Station, typ: str, coaches: Optional[int], wagons: Optional[int]) -> list[str]:
    """§46/§47.3: a PF is allowed when the stock fits and its use permits the type."""
    out = []
    for p in station.platforms:
        use = p.use.lower()
        if typ == "GDS":
            ok = (wagons or 0) <= p.goods_loop_wagons and p.goods_loop_wagons > 0
        elif typ == "MEMU":
            ok = (coaches or 0) <= p.length_coaches and "memu" in use
        else:
            ok = (coaches or 0) <= p.length_coaches and "memu only" not in use
        if ok:
            out.append(p.id)
    return out


NUM_WORDS = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten", "eleven", "twelve",
             "thirteen", "fourteen", "fifteen", "sixteen", "seventeen", "eighteen", "nineteen", "twenty"]


def default_aliases(pid: str) -> list[str]:
    """The alias pattern of the shipped station (EN · हिं · मर), for a platform ID like P7."""
    n = int(pid[1:])
    out = [str(n), f"pf{n}", f"pf {n}", f"pf-{n}", f"platform {n}", f"platform no. {n}"]
    if n < len(NUM_WORDS):
        out.append(f"platform {NUM_WORDS[n]}")
    return out + [f"प्लेटफॉर्म {n}", f"प्लेटफॉर्म नंबर {n}", f"प्लॅटफॉर्म {n}", f"फलाट {n}", f"पीएफ {n}"]


def with_default_aliases(station: Station) -> Station:
    """Platforms that list no names get the standard aliases, so a custom station parses like the shipped one."""
    if all(p.aliases for p in station.platforms):
        return station
    return station.model_copy(update={"platforms": [
        p if p.aliases else p.model_copy(update={"aliases": default_aliases(p.id)}) for p in station.platforms]})


def _merge_aliases(*lists: list[str]) -> list[str]:
    out, seen = [], set()
    for a in (x for lst in lists for x in lst):
        k = " ".join(a.lower().split())
        if k and k not in seen:
            seen.add(k)
            out.append(a.strip())
    return out[:40]


def _station_xlsx(wb, base: Station, settings: dict[str, int]) -> tuple[Station, list[str]]:
    """The "Station" sheet's platform table → a station; no table, or the same platforms as `base` → `base`."""
    ws = next((w for w in wb.worksheets if w.title.strip().lower() == "station"), None)
    if ws is None:
        return base.model_copy(update=settings), []
    name, code = "", ""
    head: Optional[dict[str, int]] = None
    plats: list[dict[str, Any]] = []
    errors: list[str] = []
    for n, r in enumerate(ws.iter_rows(min_row=1, max_row=60, max_col=8, values_only=True), start=1):
        cells = list(r)
        first = _norm(cells[0] if cells else "")
        if head is None:
            if n == 1 and re.match(r"^station\s*[—–-]", str(cells[0] or "").strip(), re.I):
                name = re.split(r"[—–-]", str(cells[0]).strip(), maxsplit=1)[1].strip()
            elif first == "station code" and len(cells) > 1:
                code = str(cells[1] or "").strip()
            elif first in ("pf", "platform") and any("length" in _norm(c) for c in cells):
                head = {_norm(c): i for i, c in enumerate(cells) if c is not None}
            continue
        if all(c is None or str(c).strip() == "" for c in cells):
            break                                                       # the table ends at its first blank row

        hd = head

        def col(*keys: str) -> Any:
            i = next((hd[k] for k in hd if any(k.startswith(x) for x in keys)), None)
            return cells[i] if i is not None and i < len(cells) else None
        try:
            pid = _pf(cells[0], f"Station sheet · row {n} · PF")
            length = _int(col("length"), f"Station sheet · row {n} · Length (coaches)")
            loop = _int(col("goods loop"), f"Station sheet · row {n} · Goods loop (wagons)") or 0
        except ValueError as e:
            errors.append(str(e))
            continue
        if pid is None or length is None:
            errors.append(f"Station sheet · row {n}: PF and Length (coaches) are required")
            continue
        names = [x for x in str(col("names") or "").split(",") if x.strip()]
        plats.append({"id": pid, "length_coaches": length, "goods_loop_wagons": loop,
                      "use": str(col("use") or "").strip()[:80], "names": names})
    if errors or not plats:
        return base.model_copy(update=settings), errors
    same = [(p.id, p.length_coaches, p.goods_loop_wagons, p.use) for p in base.platforms] == \
           [(p["id"], p["length_coaches"], p["goods_loop_wagons"], p["use"]) for p in plats]
    if same:
        return base.model_copy(update=settings), []
    sid = re.sub(r"[^A-Z0-9]", "", code.upper())[:8]
    sid = sid if len(sid) >= 2 else "CUSTOM"
    try:
        st = Station(id=sid, code=sid, name=(name or code or "Custom station")[:80],
                     min_gap=settings.get("min_gap", base.min_gap), max_hold=settings.get("max_hold", base.max_hold),
                     platforms=[Platform(id=p["id"], name=f"PF {p['id'][1:]}", length_coaches=p["length_coaches"],
                                         goods_loop_wagons=p["goods_loop_wagons"], use=p["use"],
                                         aliases=_merge_aliases(p["names"], default_aliases(p["id"])))
                                for p in plats])
    except ValidationError as e:
        return base, [f"Station sheet: {er['msg'].removeprefix('Value error, ')}" for er in e.errors()]
    return st, []


def _runs_on(days: str, day: str) -> bool:
    d = days.strip()
    if not d or d.lower() == "daily":
        return True
    return day.lower()[:3] in {x.strip().lower()[:3] for x in re.split(r"[,\s/]+", d) if x.strip()}


def _rows_to_trains(rows: list[tuple], hmap: dict[str, int], station: Station, first_row: int,
                    day: str) -> tuple[list[TrainIn], list[str], list[str]]:
    trains, errors, excluded = [], [], []

    def cell(r: tuple, h: str) -> Any:
        i = hmap.get(h)
        return r[i] if i is not None and i < len(r) else None

    for n, r in enumerate(rows, start=first_row):
        if not r or all(v is None or str(v).strip() == "" for v in r):
            if trains or errors:
                break                                                   # the table ends at its first blank row
            continue
        tid = str(cell(r, "train no.") or "").strip()
        if not re.fullmatch(r"\d{5}|GDS-\d{4}", tid):
            errors.append(f"Row {n} · Train No.: {tid or '(blank)'} is not a 5-digit number or GDS-nnnn")
            continue
        try:
            typ = str(cell(r, "type") or "").strip().upper()
            coaches = _int(cell(r, "coaches"), f"Row {n} · Coaches")
            wagons = _int(cell(r, "wagons"), f"Row {n} · Wagons")
            raw_allowed = str(cell(r, "allowed pfs") or "").replace(";", ",").replace(" ", ",")
            allowed = [_pf(x, f"Row {n} · Allowed PFs") for x in raw_allowed.split(",") if x.strip()]
            days = str(cell(r, "days of run") or "Daily").strip()
            t = TrainIn(
                id=tid, name=str(cell(r, "train name") or "").strip(), type=typ,
                origin=str(cell(r, "from") or "").strip().upper(), destination=str(cell(r, "to") or "").strip().upper(),
                direction=str(cell(r, "dir") or "").strip().upper(),
                arr=_hhmm(cell(r, "arr"), f"Row {n} · Arr"), dep=_hhmm(cell(r, "dep"), f"Row {n} · Dep"),
                platform=_pf(cell(r, "sch. pf"), f"Row {n} · Sch. PF"),
                required_platform=_pf(cell(r, "req. pf"), f"Row {n} · Req. PF"),
                allowed_platforms=[a for a in allowed if a] or derive_allowed(station, typ, coaches, wagons),
                coaches=coaches, wagons=wagons, days=days,
                priority=_int(cell(r, "priority"), f"Row {n} · Priority") or 0,
            )
        except ValidationError as e:
            for err in e.errors():
                loc = ".".join(str(x) for x in err["loc"]) or "row"
                errors.append(f"Row {n} · {loc}: {err['msg'].removeprefix('Value error, ')}")
            continue
        except ValueError as e:
            errors.append(str(e))
            continue
        if not _runs_on(days, day):
            excluded.append(f"{tid} does not run on {day} ({days})")
            continue
        for msg in referential_errors(station, TimetableIn.model_construct(trains=[t])):
            errors.append(f"Row {n} · {msg}")
        trains.append(t)
    return trains, errors, excluded


def _station_settings_xlsx(wb) -> dict[str, int]:
    out: dict[str, int] = {}
    for name, field in NAMED.items():                                # named cells first (§47.4)
        try:
            dn = wb.defined_names.get(name)
            if dn is None:
                continue
            for sheet, ref in dn.destinations:
                v = wb[sheet][ref.replace("$", "")].value
                if isinstance(v, (int, float)):
                    out[field] = int(v)
        except Exception:
            continue
    for ws in wb.worksheets:                                         # then row labels
        if ws.title.strip().lower() != "station":
            continue
        for r in ws.iter_rows(min_row=1, max_row=40, max_col=3, values_only=True):
            key = _norm(r[0] if r else "")
            if key in SETTING_LABELS and len(r) > 1 and isinstance(r[1], (int, float)):
                out.setdefault(SETTING_LABELS[key], int(r[1]))
    return out


def parse_upload(filename: str, data: bytes, station: Station, sheet: Optional[str] = None,
                 day: str = "Sun") -> dict[str, Any]:
    """→ {timetable, station, sha256, sheet, settings, excluded} | {needs_sheet, sha256} | {errors, sha256}.

    `station` is the station the timetable was checked against: the loaded one (with any sheet settings),
    or the workbook's / JSON file's own station.
    """
    sha = hashlib.sha256(data).hexdigest()
    name = filename.lower().strip()
    if len(data) > MAX_BYTES:
        return {"errors": ["File is larger than 1 MB"], "sha256": sha}
    if name.endswith(".xlsm") or name.endswith(".xls"):
        return {"errors": ["Macro-enabled or legacy workbooks are refused; save as .xlsx"], "sha256": sha}
    if name.endswith(".json"):
        try:
            obj = json.loads(data.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            return {"errors": ["Not valid UTF-8 JSON"], "sha256": sha}
        try:
            if isinstance(obj, dict) and "station" in obj and isinstance(obj.get("timetable"), dict):
                station = with_default_aliases(Station.model_validate(obj["station"]))   # {station, timetable}
                tti = TimetableIn.model_validate({**obj["timetable"], "station": station.id})
            else:
                tti = TimetableIn.model_validate(obj)
        except ValidationError as e:
            return {"errors": [f"{'.'.join(str(x) for x in er['loc'])}: {er['msg']}" for er in e.errors()], "sha256": sha}
        errs = referential_errors(station, tti)
        return {"errors": errs, "sha256": sha} if errs else \
            {"timetable": tti, "station": station, "sha256": sha, "sheet": None, "settings": {}, "excluded": []}
    settings: dict[str, int] = {}
    chosen: Optional[str] = None
    st = station
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            return {"errors": ["Not a readable .xlsx workbook"], "sha256": sha}
        candidates: dict[str, dict[str, int]] = {}
        for ws in wb.worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            hm = _header_map(tuple(first))
            if hm:
                candidates[ws.title] = hm
        if not candidates:
            return {"errors": ["No sheet has the expected timetable headers: " + ", ".join(REQUIRED)], "sha256": sha}
        if sheet is None and len(candidates) > 1:
            return {"needs_sheet": list(candidates), "sha256": sha}
        chosen = sheet or next(iter(candidates))
        if chosen not in candidates:
            return {"errors": [f"Sheet {chosen!r} does not have the expected headers"], "sha256": sha}
        ws = wb[chosen]
        if (ws.max_column or 0) > MAX_COLS:
            return {"errors": [f"Sheet has more than {MAX_COLS} columns"], "sha256": sha}
        rows = [tuple(r) for r in ws.iter_rows(min_row=2, max_row=MAX_ROWS + 2, values_only=True)]
        hmap = candidates[chosen]
        wb_full = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
        settings = _station_settings_xlsx(wb_full)
        st, st_errors = _station_xlsx(wb_full, station, settings)
        if st_errors:
            return {"errors": st_errors, "sha256": sha, "sheet": chosen}
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            return {"errors": ["CSV must be UTF-8"], "sha256": sha}
        all_rows = list(csv.reader(io.StringIO(text)))
        hmap = _header_map(tuple(all_rows[0])) if all_rows else None
        if not hmap:
            return {"errors": ["First row must carry the timetable headers: " + ", ".join(REQUIRED)], "sha256": sha}
        rows = [tuple(r) for r in all_rows[1:]]
    else:
        return {"errors": ["Only .xlsx, .csv and .json files are accepted"], "sha256": sha}
    if len([r for r in rows if r and any(v not in (None, "") for v in r)]) > MAX_ROWS:
        return {"errors": [f"More than {MAX_ROWS} rows"], "sha256": sha}
    trains, errors, excluded = _rows_to_trains(rows, hmap, st, 2, day)
    if not errors and not trains:
        errors = ["No train rows found"]
    if not errors:
        try:
            tti = TimetableIn(station=st.id, day=day, trains=trains)
        except ValidationError as e:
            errors = [er["msg"].removeprefix("Value error, ") for er in e.errors()]
    if errors:
        return {"errors": errors, "sha256": sha, "sheet": chosen}
    return {"timetable": tti, "station": st, "sha256": sha, "sheet": chosen, "settings": settings, "excluded": excluded}


def to_json(tti: TimetableIn) -> dict[str, Any]:
    return json.loads(tti.model_dump_json())
