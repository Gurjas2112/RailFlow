"""Blank-but-working upload templates for your own station (SPEC §17 load {station, timetable}, §40).

    python tools/make_template.py   →  web/public/templates/railflow_template.xlsx
                                       web/public/templates/railflow_template.json

A fictional 3-platform station (Sahyadri Road, SHYR) with 8 trains and two planted conflicts, so a first upload
shows "Check as scheduled" finding them. The JSON twin is produced by importing the workbook, so the two always agree.
"""
from __future__ import annotations

import io
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from railflow.scenario_store import load_station  # noqa: E402
from railflow.timetable_import import parse_upload, to_json  # noqa: E402

OUT = ROOT / "web" / "public" / "templates"
HEAD = ["Train No.", "Train Name", "Type", "From", "To", "Dir", "Arr", "Dep", "Halt (min)", "Sch. PF", "Req. PF",
        "Allowed PFs", "Coaches", "Wagons", "Days of run", "Priority"]
PLATFORMS = [  # PF · Length (coaches) · Goods loop (wagons) · Use · Names RailFlow understands (optional)
    (1, 24, 0, "main line", ""),
    (2, 24, 50, "main line + goods loop", ""),
    (3, 12, 0, "bay · MEMU only", "bay platform"),
]
TRAINS = [
    ("12601", "Sahyadri Superfast (fictional)", "SF", "KRJ", "LNL", "DN", "10:00", "10:05", 5, 1, None, "1,2", 22, None, "Daily", 1),
    ("11605", "Deccan Mail (fictional)", "MEX", "LNL", "KRJ", "UP", "10:08", "10:18", 10, 1, None, "1,2", 22, None, "Daily", 2),
    ("51607", "Ghat Passenger (fictional)", "PAS", "KRJ", "LNL", "DN", "10:20", "10:30", 10, 2, None, "1,2", 16, None, "Daily", 3),
    ("61609", "Valley MEMU (fictional)", "MEMU", "LNL", "KRJ", "UP", "10:25", "10:35", 10, 3, 3, None, 12, None, "Daily", 3),
    ("GDS-2601", "Cement rake (fictional)", "GDS", "KRJ", "LNL", "DN", "10:40", "11:10", 30, 2, None, None, None, 45, "Daily", 4),
    ("12611", "Plateau Express (fictional)", "SF", "LNL", "KRJ", "UP", "10:50", "10:55", 5, 1, None, "1,2", 24, None, "Daily", 1),
    ("61613", "Valley MEMU (fictional)", "MEMU", "KRJ", "LNL", "DN", "11:00", "11:10", 10, 3, 3, None, 12, None, "Daily", 3),
    ("11615", "Coast Express (fictional)", "MEX", "KRJ", "LNL", "DN", "11:05", "11:15", 10, 2, None, "1,2", 22, None, "Daily", 2),
]
README = [
    "RailFlow upload template — replace the example rows with your own station and trains (fictional data only).",
    "",
    "Sheet 'Station': station name in A1 after 'Station —', the station code, minimum platform clearance and maximum hold,",
    "then one row per platform: PF number, Length (coaches), Goods loop (wagons, 0 = none), Use, optional extra names.",
    "Use containing 'MEMU only' reserves a platform for MEMUs; 'MEMU' in Use admits MEMUs. Up to 20 platforms.",
    "",
    "Sheet 'Timetable': keep the header row exactly. One row per train, times as HH:MM (00:00–23:59, no midnight wrap).",
    "Train No.: 5 digits, or GDS-nnnn for goods. Type: SF, MEX, PAS, MEMU or GDS. Dir: UP or DN. Priority: 1 (highest) to 4.",
    "Sch. PF: the timetabled platform (blank = unassigned). Req. PF: a platform the train must use (blank = any allowed).",
    "Allowed PFs: e.g. 1,2 — blank derives them from length, goods loop and use. Coaches for passenger trains, Wagons for goods.",
    "",
    "The example has two planted conflicts: PF 1 12601 → 11605 is 2 min short of the 5-min clearance, and PF 2",
    "GDS-2601 overlaps 11615. Upload it, press Check as scheduled, then Assign platforms to see them resolved.",
    "",
    "JSON alternative: railflow_template.json — {\"station\": {...}, \"timetable\": {\"trains\": [...]}}.",
    "Platform aliases are optional there; RailFlow adds 'PF 3', 'platform 3', 'प्लेटफॉर्म 3', 'फलाट 3' and the like.",
]


def build_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for line in README:
        ws.append([line])
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 120

    st = wb.create_sheet("Station")
    st.append(["Station — Sahyadri Road (fictional)"])
    st.append(["Station code", "SHYR"])
    st.append(["Minimum platform clearance (min)", 5])
    st.append(["Maximum hold (min)", 30])
    st.append(["PF", "Length (coaches)", "Goods loop (wagons)", "Use", "Names RailFlow understands (optional)"])
    for row in PLATFORMS:
        st.append(list(row))
    st["A1"].font = st["A5"].font = st["B5"].font = st["C5"].font = st["D5"].font = st["E5"].font = Font(bold=True)
    for col, w in zip("ABCDE", (34, 18, 20, 26, 40)):
        st.column_dimensions[col].width = w

    tt = wb.create_sheet("Timetable")
    tt.append(HEAD)
    for row in TRAINS:
        tt.append(list(row))
    for c in tt[1]:
        c.font = Font(bold=True)
    tt.column_dimensions["B"].width = 32
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    data = build_xlsx()
    res = parse_upload("railflow_template.xlsx", data, load_station())
    assert "timetable" in res, res.get("errors")
    station = res["station"].model_dump(mode="json", exclude={"capacity_rule", "model_routes"})
    for p in station["platforms"]:
        p.pop("aliases")                                                 # optional in JSON; filled on upload
    (OUT / "railflow_template.xlsx").write_bytes(data)
    (OUT / "railflow_template.json").write_text(
        json.dumps({"station": station, "timetable": to_json(res["timetable"])}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8")
    print(f"{res['station'].name}: {len(res['station'].platforms)} platforms, {len(res['timetable'].trains)} trains"
          f" -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
